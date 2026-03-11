from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

# Palette
NAVY = RGBColor(15, 32, 64)
TEAL = RGBColor(0, 138, 170)
SKY = RGBColor(226, 242, 248)
LIGHT = RGBColor(247, 250, 252)
DARK = RGBColor(30, 41, 59)
MUTED = RGBColor(71, 85, 105)
WHITE = RGBColor(255, 255, 255)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)


def add_title(slide, title, subtitle=None, dark=False):
    if dark:
        bg = slide.background
        bg.fill.solid()
        bg.fill.fore_color.rgb = NAVY
        tcol, scol = WHITE, RGBColor(194, 221, 255)
    else:
        bg = slide.background
        bg.fill.solid()
        bg.fill.fore_color.rgb = LIGHT
        band = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.38))
        band.fill.solid()
        band.fill.fore_color.rgb = TEAL
        band.line.fill.background()
        tcol, scol = DARK, MUTED

    tbox = slide.shapes.add_textbox(Inches(0.8), Inches(0.6), Inches(11.7), Inches(1.0))
    tf = tbox.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(38)
    p.font.bold = True
    p.font.color.rgb = tcol

    if subtitle:
        sbox = slide.shapes.add_textbox(Inches(0.8), Inches(1.65), Inches(11.6), Inches(0.7))
        sf = sbox.text_frame
        sf.text = subtitle
        sp = sf.paragraphs[0]
        sp.font.size = Pt(18)
        sp.font.color.rgb = scol


def add_card(slide, x, y, w, h, title, bullets):
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    card.fill.solid()
    card.fill.fore_color.rgb = WHITE
    card.line.color.rgb = RGBColor(216, 226, 239)

    head = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(0.45))
    head.fill.solid()
    head.fill.fore_color.rgb = SKY
    head.line.fill.background()

    ht = slide.shapes.add_textbox(Inches(x + 0.2), Inches(y + 0.07), Inches(w - 0.3), Inches(0.3))
    hp = ht.text_frame.paragraphs[0]
    hp.text = title
    hp.font.bold = True
    hp.font.size = Pt(16)
    hp.font.color.rgb = NAVY

    tx = slide.shapes.add_textbox(Inches(x + 0.25), Inches(y + 0.55), Inches(w - 0.45), Inches(h - 0.7))
    tf = tx.text_frame
    tf.word_wrap = True
    for i, b in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = b
        p.level = 0
        p.font.size = Pt(14)
        p.font.color.rgb = DARK
        p.space_after = Pt(6)


def add_section_title(slide, text):
    tb = slide.shapes.add_textbox(Inches(0.8), Inches(0.65), Inches(11.8), Inches(0.7))
    p = tb.text_frame.paragraphs[0]
    p.text = text
    p.font.size = Pt(32)
    p.font.bold = True
    p.font.color.rgb = NAVY


# Slide 1
s = prs.slides.add_slide(prs.slide_layouts[6])
add_title(
    s,
    "Stellar Dynamics Stability Prediction",
    "Research Brief based on current README.md project definition and baseline results",
    dark=True,
)

accent = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(3.0), Inches(12.0), Inches(3.8))
accent.fill.solid()
accent.fill.fore_color.rgb = RGBColor(22, 49, 90)
accent.line.fill.background()

for i, text in enumerate([
    "Goal: predict long-term orbital stability from initial conditions and early trajectories",
    "Task A: probabilistic outcome over perturbation sets",
    "Task B: early prefix to binary stable/unstable classification",
]):
    t = s.shapes.add_textbox(Inches(1.2), Inches(3.4 + i * 1.0), Inches(11.0), Inches(0.8))
    p = t.text_frame.paragraphs[0]
    p.text = text
    p.font.size = Pt(20)
    p.font.color.rgb = WHITE

# Slide 2
s = prs.slides.add_slide(prs.slide_layouts[6])
add_title(s, "Project Scope and Core Tasks", "Two complementary prediction objectives")
add_card(
    s,
    0.8,
    2.0,
    5.9,
    4.9,
    "Task A: Distribution Prediction",
    [
        "Input: initial condition with 8 perturbed trajectories",
        "Output: P(stable), P(unstable)",
        "Target from empirical count across perturbations",
        "Captures uncertainty under same physical setup",
    ],
)
add_card(
    s,
    6.8,
    2.0,
    5.7,
    4.9,
    "Task B: Prefix-to-Label Classification",
    [
        "Input: initial condition + early trajectory prefix (e.g., T=100)",
        "Output: single binary class",
        "Label rule: stable if time_unbound == 5000, else unstable",
        "Focus: early warning of long-term divergence",
    ],
)

# Slide 3
s = prs.slides.add_slide(prs.slide_layouts[6])
add_title(s, "Data Assets and Suggested Organization", "Case-level and perturbation-level structure")
add_card(
    s,
    0.8,
    2.0,
    4.2,
    4.8,
    "Initial Conditions (ICs)",
    [
        "File: dataset/ICs_sampling_0.hdf5",
        "Scale: 40,000 samples",
        "Static orbital/physical features",
        "Examples: aratios, mratios, eccs, incs, APs, LANs",
    ],
)
add_card(
    s,
    5.2,
    2.0,
    4.2,
    4.8,
    "Time-Series Trajectories",
    [
        'File: "restults_batch_sampling_0.hdf5"',
        "Scale: 8,000 trajectories x 1,001 steps",
        "Temporal evolution of orbital quantities",
        "Includes positions, velocities, energy, momentum",
    ],
)
add_card(
    s,
    9.6,
    2.0,
    2.9,
    4.8,
    "Indexing",
    [
        "case_id",
        "perturb_id (0-7)",
        "x0",
        "trajectory",
        "time_unbound",
        "diagnosis",
    ],
)

# Slide 4
s = prs.slides.add_slide(prs.slide_layouts[6])
add_title(s, "Feature Engineering Strategy", "Improving periodic feature representation")

box1 = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.9), Inches(6.0), Inches(5.0))
box1.fill.solid(); box1.fill.fore_color.rgb = WHITE
box1.line.color.rgb = RGBColor(210, 220, 232)

box2 = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.95), Inches(1.9), Inches(5.55), Inches(5.0))
box2.fill.solid(); box2.fill.fore_color.rgb = WHITE
box2.line.color.rgb = RGBColor(210, 220, 232)

l = s.shapes.add_textbox(Inches(1.1), Inches(2.2), Inches(5.4), Inches(4.3))
tf = l.text_frame
tf.word_wrap = True
left_lines = [
    "Periodic angles handled with trig transform:",
    "- APs -> sin(APs), cos(APs)",
    "- LANs -> sin(LANs), cos(LANs)",
    "",
    "Rationale:",
    "- Preserves circular continuity (0 approx 2*pi)",
    "- Prevents false distance jumps at wrap-around",
    "",
    "Inclinations:",
    "- Keep raw incs or use cos(incs)",
]
for i, line in enumerate(left_lines):
    p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
    p.text = line
    p.font.size = Pt(16 if i == 0 else 14)
    p.font.bold = i == 0
    p.font.color.rgb = DARK
    p.space_after = Pt(4)

r = s.shapes.add_textbox(Inches(7.25), Inches(2.2), Inches(4.95), Inches(4.3))
rtf = r.text_frame
rtf.word_wrap = True
for i, line in enumerate([
    "Final feature count: 15 parameters",
    "2 aratios",
    "1 mratios",
    "2 eccs",
    "2 incs or cos(incs)",
    "2 sin(APs)",
    "2 cos(APs)",
    "2 sin(LANs)",
    "2 cos(LANs)",
]):
    p = rtf.paragraphs[0] if i == 0 else rtf.add_paragraph()
    p.text = line
    p.font.size = Pt(17 if i == 0 else 14)
    p.font.bold = i == 0
    p.font.color.rgb = NAVY if i == 0 else DARK
    p.space_after = Pt(3)

# Slide 5
s = prs.slides.add_slide(prs.slide_layouts[6])
add_title(s, "Label Construction and Training Objectives", "Aligned labels for probabilistic and binary settings")
add_card(
    s,
    0.8,
    2.0,
    6.0,
    4.9,
    "Task A Labels (per case_id)",
    [
        "Aggregate 8 perturbation outcomes",
        "p_stable = n_stable / 8",
        "p_diverge = 1 - p_stable",
        "Train with KL/Cross-Entropy or soft-target regression",
    ],
)
add_card(
    s,
    7.0,
    2.0,
    5.5,
    4.9,
    "Task B Labels (per sample)",
    [
        "y = 1 if time_unbound == 5000 (stable)",
        "y = 0 otherwise (unstable)",
        "Use only trajectory prefix [0:T]",
        "T tested at 100, 200, 500",
    ],
)

# Slide 6 - Task A results
s = prs.slides.add_slide(prs.slide_layouts[6])
add_title(s, "Baseline Results: Task A", "Generalization gap indicates overfitting risk")

chart_data = CategoryChartData()
chart_data.categories = ["MLP", "Transformer", "XGBoost"]
chart_data.add_series("Test MSE", (0.1236, 0.1417, 0.0983))
chart_data.add_series("Test MAE", (0.2224, 0.2447, 0.2435))

chart = s.shapes.add_chart(
    XL_CHART_TYPE.COLUMN_CLUSTERED,
    Inches(0.8), Inches(2.0), Inches(7.0), Inches(4.8),
    chart_data,
).chart
chart.has_legend = True
chart.legend.include_in_layout = False
chart.category_axis.has_major_gridlines = False
chart.value_axis.has_major_gridlines = True

notes = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.1), Inches(2.0), Inches(4.4), Inches(4.8))
notes.fill.solid(); notes.fill.fore_color.rgb = WHITE
notes.line.color.rgb = RGBColor(210, 220, 232)

nt = s.shapes.add_textbox(Inches(8.35), Inches(2.25), Inches(3.95), Inches(4.3))
ntf = nt.text_frame
for i, line in enumerate([
    "Observed Patterns",
    "- XGBoost has strongest test MSE",
    "- MLP and Transformer degrade by epoch 200",
    "- Train-vs-test gap is substantial",
    "",
    "Implication",
    "Prioritize regularization and split strategy",
]):
    p = ntf.paragraphs[0] if i == 0 else ntf.add_paragraph()
    p.text = line
    p.font.size = Pt(16 if i in (0,5) else 14)
    p.font.bold = i in (0,5)
    p.font.color.rgb = NAVY if i in (0,5) else DARK
    p.space_after = Pt(3)

# Slide 7 - Task B results
s = prs.slides.add_slide(prs.slide_layouts[6])
add_title(s, "Baseline Results: Task B", "Performance improves with longer trajectory prefix")

chart_data = CategoryChartData()
chart_data.categories = ["T=100", "T=200"]
chart_data.add_series("Linear SVM Acc", (0.8083, 0.8420))
chart_data.add_series("1D-CNN Acc", (0.8674, 0.8886))
chart_data.add_series("Attention Acc", (0.8670, 0.8974))

chart = s.shapes.add_chart(
    XL_CHART_TYPE.COLUMN_CLUSTERED,
    Inches(0.8), Inches(2.0), Inches(7.3), Inches(4.8),
    chart_data,
).chart
chart.has_legend = True
chart.value_axis.minimum_scale = 0.75
chart.value_axis.maximum_scale = 0.96

kpi = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.35), Inches(2.0), Inches(4.0), Inches(2.0))
kpi.fill.solid(); kpi.fill.fore_color.rgb = RGBColor(235, 247, 252)
kpi.line.fill.background()
kt = s.shapes.add_textbox(Inches(8.6), Inches(2.2), Inches(3.5), Inches(1.6))
ktf = kt.text_frame
kp = ktf.paragraphs[0]
kp.text = "Best observed at T=500"
kp.font.size = Pt(14)
kp.font.bold = True
kp.font.color.rgb = NAVY
p2 = ktf.add_paragraph()
p2.text = "Attention: AUC 0.9783 | F1 0.9210 | Acc 0.9410"
p2.font.size = Pt(12)
p2.font.color.rgb = DARK
p3 = ktf.add_paragraph()
p3.text = "1D-CNN: AUC 0.9783 | F1 0.9087 | Acc 0.9308"
p3.font.size = Pt(12)
p3.font.color.rgb = DARK

obs = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.35), Inches(4.15), Inches(4.0), Inches(2.65))
obs.fill.solid(); obs.fill.fore_color.rgb = WHITE
obs.line.color.rgb = RGBColor(210, 220, 232)
ot = s.shapes.add_textbox(Inches(8.55), Inches(4.35), Inches(3.6), Inches(2.2))
otf = ot.text_frame
for i, line in enumerate([
    "Takeaways",
    "- Prefix length strongly affects outcome confidence",
    "- Deep sequence models outperform linear baselines",
    "- T=200 already gives a strong tradeoff",
]):
    p = otf.paragraphs[0] if i == 0 else otf.add_paragraph()
    p.text = line
    p.font.size = Pt(15 if i == 0 else 13)
    p.font.bold = i == 0
    p.font.color.rgb = NAVY if i == 0 else DARK
    p.space_after = Pt(3)

# Slide 8
s = prs.slides.add_slide(prs.slide_layouts[6])
add_title(s, "Recommended Next Research Steps", "Execution priorities for the next iteration")

steps = [
    ("1. Data Integrity Pass", "Confirm case/perturb mapping and fix file naming inconsistencies before training."),
    ("2. Evaluation Protocol", "Use case-level split to prevent leakage across perturbations."),
    ("3. Task A Modeling", "Calibrated probabilistic models + uncertainty-aware loss and regularization."),
    ("4. Task B Modeling", "Benchmark at T=100/200/500 with consistent metrics and ablations."),
    ("5. Deliverables", "Prepare reproducible training scripts and an updated result table by model/task."),
]
for i, (h, d) in enumerate(steps):
    y = 1.8 + i * 1.08
    tag = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(0.9), Inches(y), Inches(0.38), Inches(0.38))
    tag.fill.solid(); tag.fill.fore_color.rgb = TEAL
    tag.line.fill.background()

    t = s.shapes.add_textbox(Inches(1.4), Inches(y - 0.05), Inches(11.6), Inches(0.9))
    tf = t.text_frame
    p1 = tf.paragraphs[0]
    p1.text = h
    p1.font.bold = True
    p1.font.size = Pt(17)
    p1.font.color.rgb = NAVY

    p2 = tf.add_paragraph()
    p2.text = d
    p2.font.size = Pt(13)
    p2.font.color.rgb = DARK

# Slide 9
s = prs.slides.add_slide(prs.slide_layouts[6])
add_title(s, "Summary", "Current status and decision points", dark=True)

summary = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(2.0), Inches(11.3), Inches(4.6))
summary.fill.solid(); summary.fill.fore_color.rgb = RGBColor(22, 49, 90)
summary.line.fill.background()

text = s.shapes.add_textbox(Inches(1.4), Inches(2.4), Inches(10.5), Inches(3.8))
tf = text.text_frame
points = [
    "The project is well-scoped around two useful prediction tasks: distributional uncertainty and early classification.",
    "Existing baselines already show strong signal, especially for sequence models in Task B.",
    "Main bottleneck is robust generalization and reproducible evaluation at the case level.",
    "Next milestone: standardized data pipeline + regularized models + consolidated benchmark report.",
]
for i, line in enumerate(points):
    p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
    p.text = line
    p.font.size = Pt(18)
    p.font.color.rgb = WHITE
    p.space_after = Pt(10)

out = "stellar_dynamics_research_brief.pptx"
prs.save(out)
print(out)
